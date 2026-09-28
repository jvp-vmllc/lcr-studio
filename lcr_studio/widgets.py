"""Reusable UI building blocks."""
from __future__ import annotations

import csv
import time

import numpy as np
import pyqtgraph as pg
from PySide6.QtCore import Qt, Signal
from PySide6.QtWidgets import (QButtonGroup, QFileDialog, QFrame, QGridLayout, QHBoxLayout, QLabel,
                               QLineEdit, QMessageBox, QPushButton, QSizePolicy, QVBoxLayout, QWidget)

from .engmath import fmt, parse_eng
from .theme import repolish, style_plot, theme

pg.setConfigOptions(antialias=True)

REVEAL_S = 0.4      # seconds a newly measured chart point takes to grow in


class Card(QFrame):
    """Rounded surface with an optional small caps title."""

    def __init__(self, title: str = "", parent=None, spacing: int = 8):
        super().__init__(parent)
        self.setObjectName("Card")
        self.outer = QVBoxLayout(self)
        self.outer.setContentsMargins(14, 12, 14, 14)
        self.outer.setSpacing(spacing)
        self.header = QHBoxLayout()
        self.header.setSpacing(6)
        if title:
            self.title = QLabel(title.upper())
            self.title.setObjectName("CardTitle")
            self.header.addWidget(self.title)
            self.header.addStretch(1)
            self.outer.addLayout(self.header)
        self.body = self.outer


class Segmented(QWidget):
    """Row of mutually exclusive toggle buttons. Emits changed(value) on user clicks only."""

    changed = Signal(object)

    def __init__(self, options, parent=None, columns: int | None = None):
        super().__init__(parent)
        lay = QGridLayout(self)
        lay.setContentsMargins(0, 0, 0, 0)
        lay.setSpacing(4)
        self.group = QButtonGroup(self)
        self.group.setExclusive(True)
        self.buttons = {}
        columns = columns or len(options)
        for i, opt in enumerate(options):
            value, label = opt if isinstance(opt, tuple) else (opt, str(opt))
            b = QPushButton(label)
            b.setObjectName("Seg")
            b.setCheckable(True)
            b.setCursor(Qt.PointingHandCursor)
            b.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Fixed)
            b.clicked.connect(lambda _=False, v=value: self.changed.emit(v))
            self.group.addButton(b)
            self.buttons[value] = b
            lay.addWidget(b, i // columns, i % columns)

    def set_value(self, value):
        b = self.buttons.get(value)
        if b is not None:
            b.setChecked(True)
        else:
            self.group.setExclusive(False)
            for btn in self.buttons.values():
                btn.setChecked(False)
            self.group.setExclusive(True)

    def value(self):
        for v, b in self.buttons.items():
            if b.isChecked():
                return v
        return None


class Badge(QLabel):
    def __init__(self, text: str = "", kind: str = "", parent=None):
        super().__init__(text, parent)
        self.setObjectName("Badge")
        self.set_kind(kind)

    def set_kind(self, kind: str):
        if self.property("kind") != kind:
            self.setProperty("kind", kind)
            repolish(self)


def field_label(text: str) -> QLabel:
    lab = QLabel(text)
    lab.setObjectName("FieldLabel")
    return lab


def muted(text: str = "") -> QLabel:
    lab = QLabel(text)
    lab.setObjectName("Muted")
    lab.setWordWrap(True)
    return lab


class EngEdit(QLineEdit):
    """Line edit that accepts engineering notation (4.7u, 10k, 2.2nF...)."""

    valueChanged = Signal(object)

    def __init__(self, placeholder: str = "", unit: str = "", parent=None):
        super().__init__(parent)
        self.unit = unit
        self.setPlaceholderText(placeholder)
        self.textChanged.connect(self._changed)

    def _changed(self):
        v = self.value()
        bad = bool(self.text().strip()) and v is None
        if bool(self.property("invalid")) != bad:
            self.setProperty("invalid", bad)
            repolish(self)
        self.valueChanged.emit(v)

    def value(self):
        try:
            return parse_eng(self.text()) if self.text().strip() else None
        except ValueError:
            return None

    def set_value(self, v, digits: int = 5):
        self.setText("" if v is None else fmt(v, self.unit, digits))


class RollAxis(pg.AxisItem):
    """Bottom axis of a rolling strip chart: ticks sit at round numbers of seconds before `now` (0 = right edge)."""

    def __init__(self):
        super().__init__(orientation="bottom")
        self.now = 0.0

    def tickSpacing(self, minVal, maxVal, size):
        return [(sp, (self.now % sp) if sp else off) for sp, off in super().tickSpacing(minVal, maxVal, size)]

    def tickStrings(self, values, scale, spacing):
        return [f"{round((v - self.now) * scale, 6) + 0.0:g}" for v in values]


def make_plot(title_left: str = "", units_left: str = "", title_bottom: str = "", units_bottom: str = "",
              log_x: bool = False, bottom_axis: pg.AxisItem | None = None) -> pg.PlotWidget:
    w = pg.PlotWidget(axisItems={"bottom": bottom_axis} if bottom_axis else None)
    p = w.getPlotItem()
    p.setLabel("left", title_left, units=units_left)
    p.setLabel("bottom", title_bottom, units=units_bottom)
    p.setLogMode(x=log_x, y=False)
    for ax in ("left", "bottom"):
        p.getAxis(ax).enableAutoSIPrefix(True)
    p.getViewBox().setMouseMode(pg.ViewBox.RectMode)

    def restyle(c, w=w, p=p):
        w.setBackground(c["surface"])
        style_plot(p, c)
    restyle(theme.c)
    theme.changed.connect(restyle)
    w.setMinimumHeight(140)
    return w


def smooth_curve(x, y, log_x: bool = False, n: int = 160):
    """Dense monotone-cubic curve through (x, y): smooth between points, never overshooting them."""
    xa = np.asarray(x, dtype=float)
    ya = np.asarray(y, dtype=float)
    if len(xa) < 3:
        return xa, ya
    order = np.argsort(xa)
    xa, ya = xa[order], ya[order]
    xi = np.log10(xa) if log_x else xa
    h = np.diff(xi)
    if np.any(h <= 0) or not np.all(np.isfinite(ya)):
        return xa, ya
    d = np.diff(ya) / h
    m = np.empty_like(ya)
    m[0], m[-1] = d[0], d[-1]
    for i in range(1, len(xa) - 1):                 # Fritsch-Butland slopes keep the curve monotone
        if d[i - 1] * d[i] <= 0:
            m[i] = 0.0
        else:
            w1, w2 = 2 * h[i] + h[i - 1], h[i] + 2 * h[i - 1]
            m[i] = (w1 + w2) / (w1 / d[i - 1] + w2 / d[i])
    xs = np.unique(np.concatenate([np.linspace(xi[0], xi[-1], n), xi]))   # knots included: markers sit on the line
    k = np.clip(np.searchsorted(xi, xs, side="right") - 1, 0, len(xi) - 2)
    t = (xs - xi[k]) / h[k]
    t2, t3 = t * t, t * t * t
    ys = ((2 * t3 - 3 * t2 + 1) * ya[k] + (t3 - 2 * t2 + t) * h[k] * m[k]
          + (-2 * t3 + 3 * t2) * ya[k + 1] + (t3 - t2) * h[k] * m[k + 1])
    return (10 ** xs if log_x else xs), ys


def reveal_fraction(t0: float):
    """Ease-out progress of a grow-in animation started at monotonic time t0: (0..1, finished)."""
    f = min(1.0, (time.monotonic() - t0) / REVEAL_S)
    return 1 - (1 - f) ** 3, f >= 1.0


def plot_series(plot, x, y, color, name=None, log_x: bool = False, style=Qt.SolidLine, width: int = 2,
                symbol_size: int = 7, reveal: float = 1.0):
    """Smooth line through the points plus a marker on every measured point. Returns the line item.

    reveal < 1 draws the newest point only part of the way from the previous one (grow-in animation).
    """
    x, y = list(x), list(y)
    tail = None
    if reveal < 1.0 and len(x) >= 2:
        if log_x:
            x[-1] = 10 ** (np.log10(x[-2]) + (np.log10(x[-1]) - np.log10(x[-2])) * reveal)
        else:
            x[-1] = x[-2] + (x[-1] - x[-2]) * reveal
        y[-1] = y[-2] + (y[-1] - y[-2]) * reveal
        tail = (x.pop(), y.pop())
    elif reveal < 1.0 and len(x) == 1:
        tail = (x.pop(), y.pop())                 # first point of a series just grows in place
    xs, ys = smooth_curve(x + ([tail[0]] if tail else []), y + ([tail[1]] if tail else []), log_x)
    line = plot.plot(xs, ys, pen=pg.mkPen(color, width=width, style=style), name=name)
    if x:
        plot.plot(x, y, pen=None, symbol="o", symbolSize=symbol_size, symbolBrush=color, symbolPen=None)
    if tail:
        plot.plot([tail[0]], [tail[1]], pen=None, symbol="o", symbolSize=max(1.0, symbol_size * reveal),
                  symbolBrush=color, symbolPen=None)
    return line


def ask(parent, title: str, text: str) -> bool:
    return QMessageBox.question(parent, title, text, QMessageBox.Yes | QMessageBox.No,
                                QMessageBox.No) == QMessageBox.Yes


def export_table(parent, headers, rows, default_name: str):
    """Ask for a path and save rows as CSV or Excel."""
    path, _ = QFileDialog.getSaveFileName(parent, "Export data", default_name,
                                          "Excel workbook (*.xlsx);;CSV (*.csv)")
    if not path:
        return None
    try:
        if path.lower().endswith(".xlsx"):
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill
            wb = Workbook()
            ws = wb.active
            ws.append(list(headers))
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="2563EB")
            for r in rows:
                ws.append(list(r))
            for i, h in enumerate(headers, 1):
                ws.column_dimensions[ws.cell(1, i).column_letter].width = max(11, len(str(h)) + 3)
            ws.freeze_panes = "A2"
            wb.save(path)
        else:
            with open(path, "w", newline="", encoding="utf-8-sig") as f:
                w = csv.writer(f)
                w.writerow(headers)
                w.writerows(rows)
    except OSError as exc:
        QMessageBox.warning(parent, "Export failed", str(exc))
        return None
    return path
