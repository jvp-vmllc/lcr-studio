"""Flyback transformer test: Lp / Llk measurements and a one-page report."""
from __future__ import annotations

import json
import math
import platform
import time
from datetime import datetime

import pyqtgraph as pg
from PySide6.QtCore import QSettings, Qt, QTimer
from PySide6.QtGui import QColor, QPixmap
from PySide6.QtWidgets import (QCheckBox, QComboBox, QDialog, QDoubleSpinBox, QFileDialog, QGridLayout, QHBoxLayout,
                               QHeaderView, QLabel, QLineEdit, QMessageBox, QPushButton, QScrollArea,
                               QSpinBox, QSplitter, QTableWidget, QTableWidgetItem, QVBoxLayout, QWidget)

from .. import __version__
from ..engmath import fmt
from ..flyback import FlybackProfile, build_steps, evaluate, hz_label, run_step, v_label
from ..report import loss_figures
from ..report import LEVEL_COLORS, default_meta, export_pdf, render_image
from ..theme import theme
from ..ut622e import FREQ_HZ, FREQUENCIES, LEVELS
from ..widgets import (Badge, BusyDialog, Card, EngAxis, Segmented, SmoothRange, UnitEdit, field_label, make_plot, muted,
                       plot_series, range_with_limits, reveal_fraction, watch_manual_zoom)

STATUS_KIND = {"done": "good", "running": "accent", "aborted": "warn", "error": "bad", "pending": ""}


# ------------------------------------------------------------ report preview --

class ReportPreview(QDialog):
    def __init__(self, image, on_export, parent=None):
        super().__init__(parent)
        self.setWindowTitle("Report preview — US Letter")
        self.resize(900, 1000)
        lay = QVBoxLayout(self)
        scroll = QScrollArea()
        scroll.setWidgetResizable(False)
        scroll.setAlignment(Qt.AlignHCenter)
        lab = QLabel()
        lab.setPixmap(QPixmap.fromImage(image))
        lab.setStyleSheet("background: white; border: 1px solid #ccd;")
        scroll.setWidget(lab)
        lay.addWidget(scroll, 1)
        row = QHBoxLayout()
        row.addStretch(1)
        exp = QPushButton("Export PDF…")
        exp.setObjectName("Accent")
        exp.clicked.connect(lambda: (on_export(), self.accept()))
        close = QPushButton("Close")
        close.clicked.connect(self.reject)
        row.addWidget(exp)
        row.addWidget(close)
        lay.addLayout(row)


# ------------------------------------------------------------------ panel --

class FlybackPanel(QWidget):
    TAG = "flyback"

    def __init__(self, worker, settings: QSettings | None = None, parent=None):
        super().__init__(parent)
        self.worker = worker
        self.qs = settings
        self.results: dict = {}
        self.status: dict = {}
        self.running_key = None
        self.connected = False
        self.meta = {"meter": "—", "speed": "?", "correction": "—"}
        self._loading = False
        self.busy = None                      # pop-up shown while a step runs
        self.queue: list[str] = []            # ticked steps still to run after the current one
        self._anim_t0 = None
        self._anim = QTimer(self)             # redraws while the newest point grows in
        self._anim.setInterval(16)
        self._anim.timeout.connect(self._refresh_results)

        root = QHBoxLayout(self)
        root.setContentsMargins(12, 12, 12, 12)
        root.setSpacing(10)
        root.addWidget(self._build_setup())
        split = QSplitter(Qt.Vertical)
        split.setChildrenCollapsible(False)
        split.addWidget(self._build_steps())
        split.addWidget(self._build_results())
        split.setSizes([250, 600])
        root.addWidget(split, 1)

        self._load_saved_profile()
        self._rebuild_steps()
        self.set_connected(False)

    # ============================================================ setup ==
    def _build_setup(self):
        scroll = QScrollArea()
        scroll.setWidgetResizable(True)
        scroll.setHorizontalScrollBarPolicy(Qt.ScrollBarAlwaysOff)
        scroll.setFixedWidth(356)
        inner = QWidget()
        scroll.setWidget(inner)
        self.setup_scroll = scroll
        col = QVBoxLayout(inner)
        col.setContentsMargins(0, 0, 8, 0)
        col.setSpacing(10)

        unit = Card("Unit under test")
        g = QGridLayout()
        g.setVerticalSpacing(6)
        self.part = QLineEdit()
        self.part.setPlaceholderText("e.g. FBT-EE25-12V")
        self.station = QLineEdit(platform.node())
        self.station.setReadOnly(True)
        self.station.setToolTip("This computer's name — printed on the report as the test station")
        self.notes = QLineEdit()
        self.notes.setPlaceholderText("printed on the report")
        for i, (lab, wdg) in enumerate([("Part number", self.part), ("Station", self.station),
                                        ("Notes", self.notes)]):
            g.addWidget(field_label(lab), i, 0)
            g.addWidget(wdg, i, 1)
        unit.body.addLayout(g)
        row = QHBoxLayout()
        load = QPushButton("Load profile…")
        load.clicked.connect(self._load_profile_file)
        save = QPushButton("Save profile…")
        save.clicked.connect(self._save_profile_file)
        row.addWidget(load)
        row.addWidget(save)
        unit.body.addLayout(row)
        col.addWidget(unit)

        def cond_pickers():
            fq = QComboBox()
            for freq in FREQUENCIES:
                fq.addItem(hz_label(freq), freq)
            lv = QComboBox()
            for level in LEVELS:
                lv.addItem(v_label(level), level)
            return fq, lv

        def spec_card(title, rows):
            card = Card(title)
            g = QGridLayout()
            g.setVerticalSpacing(6)
            for i, (lab, wdg) in enumerate(rows):
                g.addWidget(field_label(lab), i, 0)
                g.addWidget(wdg, i, 1)
            card.body.addLayout(g)
            return card

        self.lp_freq, self.lp_level = cond_pickers()
        self.lp_nom = UnitEdit("H", required=True)
        self.lp_tol = QDoubleSpinBox()
        self.lp_tol.setRange(0.1, 50)
        self.lp_tol.setSuffix(" %")
        self.lp_tol.setValue(10)
        self.lp_equ = Segmented([("SER", "Series"), ("PAR", "Parallel")])
        col.addWidget(spec_card("Lp specification", [("Test frequency", self.lp_freq), ("Test level", self.lp_level),
                                                     ("Nominal", self.lp_nom), ("Tolerance ±", self.lp_tol),
                                                     ("Circuit model", self.lp_equ)]))

        self.llk_freq, self.llk_level = cond_pickers()
        self.llk_max = UnitEdit("H", required=True)
        self.llk_pct = UnitEdit("%", prefixes=("",), required=True)
        self.llk_equ = Segmented([("SER", "Series"), ("PAR", "Parallel")])
        llk = spec_card("Llk specification", [("Test frequency", self.llk_freq), ("Test level", self.llk_level),
                                              ("Llk max", self.llk_max), ("Llk / Lp max", self.llk_pct),
                                              ("Circuit model", self.llk_equ)])
        self.spec_hint = muted("")
        llk.body.addWidget(self.spec_hint)
        col.addWidget(llk)

        mat = Card("Test matrix")
        mat.body.addWidget(field_label("Frequencies"))
        self.freq_boxes = {}
        g = QGridLayout()
        for i, f in enumerate(FREQUENCIES):
            cb = QCheckBox(hz_label(f))
            cb.setChecked(True)
            self.freq_boxes[f] = cb
            g.addWidget(cb, i // 3, i % 3)
        mat.body.addLayout(g)
        mat.body.addWidget(field_label("Test levels"))
        self.level_boxes = {}
        row = QHBoxLayout()
        for lv in LEVELS:
            cb = QCheckBox(v_label(lv))
            cb.setChecked(True)
            self.level_boxes[lv] = cb
            row.addWidget(cb)
        mat.body.addLayout(row)
        g = QGridLayout()
        self.settle = QSpinBox()
        self.settle.setRange(0, 20)
        self.settle.setValue(2)
        self.navg = QSpinBox()
        self.navg.setRange(1, 50)
        self.navg.setValue(5)
        g.addWidget(field_label("Settle (discard)"), 0, 0)
        g.addWidget(self.settle, 0, 1)
        g.addWidget(field_label("Average"), 1, 0)
        g.addWidget(self.navg, 1, 1)
        mat.body.addLayout(g)
        col.addWidget(mat)
        col.addStretch(1)

        self.part.textChanged.connect(lambda _t: self._profile_changed())
        for w in (self.lp_nom, self.llk_max, self.llk_pct):
            w.valueChanged.connect(lambda _v: self._profile_changed())
        for w in (self.lp_freq, self.lp_level, self.llk_freq, self.llk_level):
            w.currentIndexChanged.connect(lambda _i: self._profile_changed())
        for w in (self.lp_tol,):
            w.valueChanged.connect(lambda _v: self._profile_changed())
        for w in (self.settle, self.navg):
            w.valueChanged.connect(lambda _v: self._profile_changed())
        for seg in (self.lp_equ, self.llk_equ):
            seg.changed.connect(lambda _v: self._profile_changed())
        for cb in list(self.freq_boxes.values()) + list(self.level_boxes.values()):
            cb.toggled.connect(lambda _on: self._profile_changed())
        return scroll

    # ============================================================ steps ==
    def _build_steps(self):
        card = Card("Run test")
        self.step_grid = QGridLayout()
        self.step_grid.setHorizontalSpacing(14)
        self.step_grid.setVerticalSpacing(4)
        self.step_grid.setColumnStretch(2, 1)
        self.step_boxes: dict[str, QCheckBox] = {}
        self.step_status: dict[str, QLabel] = {}
        card.body.addLayout(self.step_grid)
        row = QHBoxLayout()
        self.run_btn = QPushButton("Run ticked steps")
        self.run_btn.setObjectName("Accent")
        self.run_btn.clicked.connect(self.run_ticked)
        self.abort_btn = QPushButton("Abort")
        self.abort_btn.setEnabled(False)
        self.abort_btn.clicked.connect(self.worker.abort_job)
        self.new_btn = QPushButton("New unit")
        self.new_btn.setToolTip("Clear results for the next transformer (keeps the profile)")
        self.new_btn.clicked.connect(self.new_unit)
        row.addWidget(self.run_btn, 1)
        row.addWidget(self.abort_btn)
        row.addWidget(self.new_btn)
        card.body.addLayout(row)
        self.instruction = QLabel("")
        self.instruction.setWordWrap(True)
        self.instruction.setAlignment(Qt.AlignTop | Qt.AlignLeft)
        card.body.addWidget(self.instruction)
        self.progress_text = muted("")
        card.body.addWidget(self.progress_text)
        self.hint = muted("Run open/short correction on the meter with the same leads before testing.")
        card.body.addWidget(self.hint)
        return card

    # ========================================================== results ==
    def _build_results(self):
        card = Card("Results")
        self.verdict_badge = Badge("NOT TESTED")
        card.header.addWidget(self.verdict_badge)
        prev = QPushButton("Preview report")
        prev.clicked.connect(self.preview)
        pdf = QPushButton("Export PDF…")
        pdf.setObjectName("Accent")
        pdf.clicked.connect(self.export_pdf)
        xls = QPushButton("Export data…")
        xls.clicked.connect(self.export_data)
        for b in (prev, xls, pdf):
            card.header.addWidget(b)
        body = QHBoxLayout()
        charts = QVBoxLayout()
        self.lp_plot = make_plot("Lp", "H", "Frequency", "", log_x=True, left_axis=EngAxis("H"))
        self.llk_plot = make_plot("Llk", "H", "Frequency", "", log_x=True, left_axis=EngAxis("H"))
        ticks = [[(math.log10(FREQ_HZ[f]), f.replace("Hz", "")) for f in FREQUENCIES if f != "120Hz"],
                 [(math.log10(120), "")]]
        for plot in (self.lp_plot, self.llk_plot):
            ax = plot.getPlotItem().getAxis("bottom")
            ax.enableAutoSIPrefix(False)
            ax.setTicks(ticks)
            plot.getPlotItem().setLabel("bottom", "Frequency (Hz)")
            plot.getPlotItem().setXRange(math.log10(90), math.log10(110e3), padding=0.02)
            plot.getPlotItem().addLegend(offset=(8, 4), colCount=3)
        # specification limits: Lp nominal ± tolerance as a band, Llk max as a line
        self.lp_band = pg.LinearRegionItem(orientation="horizontal", movable=False)
        self.lp_band.setZValue(-10)
        self.lp_nom_line = pg.InfiniteLine(angle=0, movable=False, label="nominal", labelOpts={"position": 0.8})
        self.llk_line = pg.InfiniteLine(angle=0, movable=False, label="Llk max", labelOpts={"position": 0.8})
        self._restyle_limits(theme.c)
        theme.changed.connect(self._restyle_limits)
        self.ease = {"lp": SmoothRange(self.lp_plot, padding=0), "llk": SmoothRange(self.llk_plot, padding=0)}
        self._logy = {"lp": False, "llk": False}
        self._follow = True
        watch_manual_zoom(self, self.lp_plot, self.llk_plot)
        charts.addWidget(self.lp_plot)
        charts.addWidget(self.llk_plot)
        body.addLayout(charts, 4)
        self.summary = QTableWidget(0, 4)
        self.summary.setHorizontalHeaderLabels(["Parameter", "Measured", "Limit", ""])
        self.summary.setWordWrap(False)
        self.summary.verticalHeader().hide()
        self.summary.setEditTriggers(QTableWidget.NoEditTriggers)
        self.summary.setAlternatingRowColors(True)
        hh = self.summary.horizontalHeader()
        hh.setSectionResizeMode(0, QHeaderView.Stretch)
        for i in (1, 2, 3):
            hh.setSectionResizeMode(i, QHeaderView.ResizeToContents)
        body.addWidget(self.summary, 5)
        card.body.addLayout(body, 1)
        return card

    # ======================================================== profile I/O ==
    def profile(self) -> FlybackProfile:
        return FlybackProfile(
            part_number=self.part.text().strip(),
            lp_freq=self.lp_freq.currentData(), lp_level=self.lp_level.currentData(),
            llk_freq=self.llk_freq.currentData(), llk_level=self.llk_level.currentData(),
            lp_nom=self.lp_nom.value(), lp_tol=self.lp_tol.value(), llk_max=self.llk_max.value(),
            llk_pct_max=self.llk_pct.value(), lp_equ=self.lp_equ.value() or "SER", llk_equ=self.llk_equ.value() or "SER",
            freqs=[f for f, cb in self.freq_boxes.items() if cb.isChecked()],
            levels=[lv for lv, cb in self.level_boxes.items() if cb.isChecked()],
            settle=self.settle.value(), navg=self.navg.value())

    def set_profile(self, p: FlybackProfile):
        self._loading = True
        try:
            self.part.setText(p.part_number)
            for box, value in ((self.lp_freq, p.lp_freq), (self.lp_level, p.lp_level),
                               (self.llk_freq, p.llk_freq), (self.llk_level, p.llk_level)):
                box.setCurrentIndex(max(0, box.findData(value)))
            self.lp_nom.set_value(p.lp_nom, 4)
            self.lp_tol.setValue(p.lp_tol)
            self.llk_max.set_value(p.llk_max, 4)
            self.llk_pct.set_value(p.llk_pct_max)
            self.lp_equ.set_value(p.lp_equ)
            self.llk_equ.set_value(p.llk_equ)
            for f, cb in self.freq_boxes.items():
                cb.setChecked(f in p.freqs)
            for lv, cb in self.level_boxes.items():
                cb.setChecked(lv in p.levels)
            self.settle.setValue(p.settle)
            self.navg.setValue(p.navg)
        finally:
            self._loading = False
        self._profile_changed()

    def _profile_changed(self):
        if self._loading:
            return
        if self.qs is not None:
            self.qs.setValue("flyback/profile", json.dumps(self.profile().to_dict()))
        self._rebuild_steps()
        self._results_changed()
        self._update_run_enabled()

    def _load_saved_profile(self):
        p = FlybackProfile()
        if self.qs is not None:
            raw = self.qs.value("flyback/profile", "")
            if raw:
                try:
                    p = FlybackProfile.from_dict(json.loads(raw))
                except (ValueError, TypeError):
                    pass
        self.set_profile(p)

    def _load_profile_file(self):
        path, _ = QFileDialog.getOpenFileName(self, "Load transformer profile", "", "Profile (*.json)")
        if path:
            try:
                with open(path, encoding="utf-8") as f:
                    self.set_profile(FlybackProfile.from_dict(json.load(f)))
            except (OSError, ValueError, TypeError) as exc:
                QMessageBox.warning(self, "Load failed", str(exc))

    def _save_profile_file(self):
        p = self.profile()
        path, _ = QFileDialog.getSaveFileName(self, "Save transformer profile", f"{p.part_number or 'flyback'}.json",
                                              "Profile (*.json)")
        if path:
            with open(path, "w", encoding="utf-8") as f:
                json.dump(p.to_dict(), f, indent=2)

    # ============================================================ steps ==
    def _rebuild_steps(self):
        self.steps = build_steps(self.profile())
        keys = {s.key for s in self.steps}
        self.results = {k: v for k, v in self.results.items() if k in keys}
        for i, s in enumerate(self.steps):
            if s.key not in self.step_boxes:
                box = QCheckBox()
                box.setChecked(True)
                box.toggled.connect(lambda _on: self._show_step())
                label = QLabel("")
                self.step_grid.addWidget(box, i, 0)
                self.step_grid.addWidget(label, i, 1)
                self.step_boxes[s.key], self.step_status[s.key] = box, label
            self.step_boxes[s.key].setText(f"{i + 1}. {s.title}")
            self.step_boxes[s.key].setEnabled(not self.running_key)
            st = self.status.get(s.key, "pending")
            kind = STATUS_KIND.get(st, "")
            color = theme.c.get(kind, theme.c["muted"]) if kind else theme.c["muted"]
            summary = self._step_summary(s.key)
            self.step_status[s.key].setText(f"<b style='color:{color}'>{st.upper()}</b>"
                                            + (f" · {summary}" if summary else ""))
        self._show_step()

    def _step_summary(self, key):
        res = self.results.get(key)
        if not res or not res["rows"]:
            return ""
        p = self.profile()
        freq, level = (p.lp_freq, p.lp_level) if key == "lp" else (p.llk_freq, p.llk_level)
        pt = next((r for r in res["rows"] if r["freq"] == freq and r["level"] == level), None)
        if pt is None:
            return f"{fmt(res['rows'][-1]['L'], 'H', 4)} (latest point)"
        return f"{fmt(pt['L'], 'H', 4)} @ {hz_label(freq)}, {v_label(level)}"

    def _ticked(self):
        return [s for s in self.steps if self.step_boxes[s.key].isChecked()]

    def _next_step(self):
        """The ticked step that runs next: the first not yet done, else the first ticked."""
        ticked = self._ticked()
        return next((s for s in ticked if self.status.get(s.key) != "done"), ticked[0] if ticked else None)

    def _show_step(self):
        s = self._next_step()
        if s is None:
            self.instruction.setText("Tick the steps to run.")
            return
        n = len(s.freqs) * len(s.levels)
        self.instruction.setText(f"<b>Next: {s.title}</b><br>{s.instruction}<br>"
                                 f"<span style='color:{theme.c['muted']}'>{n} point(s): "
                                 f"{', '.join(hz_label(f) for f in s.freqs)} × {', '.join(v_label(v) for v in s.levels)}"
                                 f" · {'series' if s.equ == 'SER' else 'parallel'} model</span>")

    def _spec_complete(self) -> bool:
        """Every limit in the specification holds a positive value."""
        return all(v is not None and v > 0 for v in (self.lp_nom.value(), self.llk_max.value(),
                                                     self.llk_pct.value()))

    def _update_run_enabled(self):
        self.run_btn.setEnabled(self.connected and not self.running_key)
        if self._spec_complete():
            self.spec_hint.setText("Each limit is checked at its own test frequency and level.")
            self.spec_hint.setStyleSheet("")
        else:
            self.spec_hint.setText("Fill in every limit to run a step.")

    def _flag_missing(self):
        """Run was pressed with limits missing: say so and jump to the first empty one."""
        self.spec_hint.setText("⚠ Fill in every limit before running a step.")
        self.spec_hint.setStyleSheet(f"color: {theme.c['bad']};")
        self.progress_text.setText("Fill in every limit in the specification first.")
        for w in (self.lp_nom, self.llk_max, self.llk_pct):
            if w.value() is None or w.value() <= 0:
                self.setup_scroll.ensureWidgetVisible(w)
                w.setFocus()
                break

    def run_ticked(self):
        """Run every ticked step in order, asking to rewire between steps."""
        if self.running_key:
            return
        if not self._spec_complete():
            self._flag_missing()
            return
        ticked = self._ticked()
        if not ticked:
            self.progress_text.setText("Tick at least one step.")
            return
        self.queue = [s.key for s in ticked[1:]]
        self.run_step(ticked[0].key)

    def run_step(self, key: str):
        step = next((s for s in self.steps if s.key == key), None)
        if step is None or self.running_key or not self._spec_complete():
            return
        p = self.profile()
        self.running_key = step.key
        self.status[step.key] = "running"
        self.results[step.key] = {"step": step.key, "rows": [], "aborted": False, "equ": step.equ}
        self._set_running(True)
        self._rebuild_steps()
        self.worker.job(lambda m, ctx, s=step: run_step(m, ctx, s, p.settle, p.navg), tag=self.TAG)
        self.busy = BusyDialog("Measuring", step.title, self.worker.abort_job, self)
        self.busy.open_centered()

    def _continue_queue(self, key):
        if self.busy is not None:
            self.busy.accept()
            self.busy = None
        self.run_step(key)

    def _stop_queue(self):
        self.queue = []
        if self.busy is not None:
            self.busy.accept()
            self.busy = None
        self.progress_text.setText("Stopped. The remaining ticked steps were not run.")

    def _set_running(self, on):
        self.abort_btn.setEnabled(on)
        self.new_btn.setEnabled(not on)
        self._update_run_enabled()

    def on_progress(self, tag, i, n, msg):
        if tag == self.TAG:
            self.progress_text.setText(msg)
            if self.busy is not None:
                detail = msg.split(" · ", 1)[1] if " · " in msg else msg
                self.busy.progress(i, n, f"{detail}   ·   {i} / {n}")

    def on_partial(self, tag, data):
        if tag == self.TAG and self.running_key and data.get("step") == self.running_key:
            self.results[self.running_key]["rows"].append(data["row"])
            self._anim_t0 = time.monotonic()
            self._anim.start()
            self._results_changed()

    def on_done(self, tag, result):
        if tag != self.TAG or not self.running_key:
            return
        key = self.running_key
        self.running_key = None
        self.results[key] = result
        aborted = bool(result.get("aborted"))
        self.status[key] = "aborted" if aborted else "done"
        self._set_running(False)
        self._rebuild_steps()
        self._results_changed()
        if aborted:
            self.queue = []
            self.progress_text.setText("Aborted. Partial data kept.")
        elif self.queue:
            key = self.queue.pop(0)
            nxt = next(s for s in self.steps if s.key == key)
            self.progress_text.setText(f"Step complete. Wire up for {nxt.title}.")
            if self.busy is not None:                # same pop-up: what to wire next, Continue or Stop
                self.busy.prompt("Next step", nxt.title, nxt.instruction,
                                 on_continue=lambda _=False, k=nxt.key: self._continue_queue(k),
                                 on_stop=self._stop_queue)
            return
        elif all(self.status.get(s.key) == "done" for s in self._ticked()):
            self.progress_text.setText("All steps complete. Export the report.")
        else:
            self.progress_text.setText("Step complete.")
        if self.busy is not None:
            self.busy.finish(self.progress_text.text(), ok=not aborted)
            self.busy = None

    def on_error(self, tag, msg):
        if tag != self.TAG or not self.running_key:
            return
        self.queue = []
        self.status[self.running_key] = "error"
        self.results.pop(self.running_key, None)
        self.running_key = None
        self._set_running(False)
        self.progress_text.setText(f"Step failed: {msg}")
        self._rebuild_steps()
        if self.busy is not None:
            self.busy.finish(f"Failed: {msg}", ok=False, delay_ms=2500)
            self.busy = None

    def new_unit(self):
        self.queue = []
        self.results, self.status = {}, {}
        self._rebuild_steps()
        self._results_changed()

    # =========================================================== results ==
    def _restyle_limits(self, c):
        self.lp_band.setBrush(pg.mkBrush(c["good"] + "22"))
        for line in self.lp_band.lines:
            line.setPen(pg.mkPen(c["good"], width=1, style=Qt.DashLine))
        self.lp_nom_line.setPen(pg.mkPen(c["muted"], width=1, style=Qt.DotLine))
        self.lp_nom_line.label.setColor(c["muted"])
        self.llk_line.setPen(pg.mkPen(c["bad"], width=1, style=Qt.DashLine))
        self.llk_line.label.setColor(c["bad"])

    def _place_limits(self, item, key, p, logy):
        """Re-add the limit band / line after a clear; returns the limit values shown, for the y range."""
        tr = math.log10 if logy else (lambda v: v)        # items sit in view coordinates: log10 in log mode
        if key == "lp":
            item.addItem(self.lp_band)
            item.addItem(self.lp_nom_line)
            on = bool(p.lp_nom)
            self.lp_band.setVisible(on)
            self.lp_nom_line.setVisible(on)
            if not on:
                return []
            lo, hi = p.lp_nom * (1 - p.lp_tol / 100), p.lp_nom * (1 + p.lp_tol / 100)
            self.lp_band.setRegion((tr(lo), tr(hi)))
            self.lp_nom_line.setPos(tr(p.lp_nom))
            return [lo, hi]
        item.addItem(self.llk_line)
        self.llk_line.setVisible(bool(p.llk_max))
        if not p.llk_max:
            return []
        self.llk_line.setPos(tr(p.llk_max))
        return [p.llk_max]

    def _results_changed(self):
        self._follow = True                      # new data or limits: follow again after a manual zoom
        self._refresh_results()

    def _refresh_results(self):
        if not hasattr(self, "summary"):
            return
        reveal = 1.0
        if self._anim_t0 is not None:
            reveal, done = reveal_fraction(self._anim_t0)
            if done:
                self._anim_t0 = None
                self._anim.stop()
        live = self.results.get(self.running_key, {}).get("rows") if self.running_key else None
        newest = live[-1] if live else None
        p = self.profile()
        for plot, key in ((self.lp_plot, "lp"), (self.llk_plot, "llk")):
            item = plot.getPlotItem()
            item.clear()
            if item.legend:
                item.legend.clear()
            res = self.results.get(key)
            ys = [r["L"] for r in res["rows"] if r["L"] > 0] if res else []
            logy = bool(ys) and max(ys) / min(ys) > 10   # a sweep spanning a decade or more reads on a log axis
            if logy != self._logy[key]:
                self._logy[key] = logy
                self.ease[key].clear()
            item.setLogMode(x=True, y=logy)
            item.getAxis("left").enableAutoSIPrefix(not logy)     # log ticks carry their own units
            item.setLabel("left", "Lp" if key == "lp" else "Llk", units="" if logy else "H")
            limits = self._place_limits(item, key, p, logy)
            freq, level = (p.lp_freq, p.lp_level) if key == "lp" else (p.llk_freq, p.llk_level)
            for lv in (LEVELS if res else ()):
                lv_rows = sorted((r for r in res["rows"] if r["level"] == lv), key=lambda r: r["hz"])
                if lv_rows:
                    rv = reveal if lv_rows[-1] is newest else 1.0
                    plot_series(plot, [r["hz"] for r in lv_rows], [r["L"] for r in lv_rows], LEVEL_COLORS[lv],
                                v_label(lv), log_x=True, symbol_size=6, reveal=rv)
            spec_pt = next((r for r in res["rows"] if r["freq"] == freq and r["level"] == level), None) if res else None
            if spec_pt is not None and spec_pt is not newest:       # cross-hair on the point the limit is checked at
                c = theme.c
                pen = pg.mkPen(c["text"], width=1, style=Qt.DashLine)
                opts = dict(color=c["text"], fill=pg.mkBrush(c["surface"] + "dd"), movable=False)
                item.addItem(pg.InfiniteLine(pos=math.log10(spec_pt["hz"]), angle=90, pen=pen,
                                             label=f"{hz_label(freq)}, {v_label(level)}",
                                             labelOpts=dict(opts, position=0.93)))
                item.addItem(pg.InfiniteLine(pos=math.log10(spec_pt["L"]) if logy else spec_pt["L"], angle=0, pen=pen,
                                             label=fmt(spec_pt["L"], "H", 4), labelOpts=dict(opts, position=0.06)))
            if ys or limits:                     # fit the curves; a limit joins the view once it is near
                tr = math.log10 if logy else (lambda v: v)
                lims = [tr(v) for v in limits if v > 0]
                vals = [tr(v) for v in ys]
                lo, hi = range_with_limits(vals, lims) if vals else (min(lims), max(lims))
                pad = (hi - lo) * 0.12 or abs(hi) * 0.05 or 1e-9
                self.ease[key].set_target(lo - pad, hi + pad)
        moving = False
        if self._follow:
            for e in self.ease.values():
                moving = e.tick() or moving
        if moving or self._anim_t0 is not None:     # keep the display timer running until everything settles
            if not self._anim.isActive():
                self._anim.start()
        else:
            self._anim.stop()
        rows, verdict = evaluate(p, self.results)
        self.summary.setRowCount(len(rows))
        c = theme.c
        for i, r in enumerate(rows):
            for j, text in enumerate([r["param"], r["value"], r["limit"], r["status"] if r["status"] != "INFO" else ""]):
                it = QTableWidgetItem(text)
                it.setToolTip(f"{r['param']} — {r['cond']}")
                if j == 3 and r["status"] in ("PASS", "FAIL"):
                    it.setForeground(QColor(c["good"] if r["status"] == "PASS" else c["bad"]))
                    f = it.font()
                    f.setBold(True)
                    it.setFont(f)
                self.summary.setItem(i, j, it)
        has_data = any(v["rows"] for v in self.results.values())
        self.verdict_badge.setText(verdict if has_data else "NOT TESTED")
        self.verdict_badge.set_kind({"PASS": "good", "FAIL": "bad", "INCOMPLETE": "warn"}.get(verdict, "accent")
                                    if has_data else "")

    # ============================================================ report ==
    def set_connected(self, on, idn: str = ""):
        self.connected = on
        if on and idn:
            parts = idn.split(",")
            self.meta["meter"] = " · ".join(x for x in (parts[1] if len(parts) > 1 else idn,
                                                           f"S/N {parts[2]}" if len(parts) > 2 else "",
                                                           f"FW {parts[3]}" if len(parts) > 3 else "") if x)
        self._set_running(bool(self.running_key))

    def apply_settings(self, st):
        self.meta["speed"] = st.get("speed") or self.meta["speed"]
        if "open_corr" in st:
            self.meta["correction"] = f"Open {'✓' if st.get('open_corr') else '✗'}   Short {'✓' if st.get('short_corr') else '✗'}"
            ok = st.get("open_corr") and st.get("short_corr")
            self.hint.setText("Open/short correction is active on the meter." if ok else
                              "⚠ Open/short correction is off. Run it on the meter with the same leads.")

    def _meta(self):
        return default_meta(station=self.station.text().strip(), notes=self.notes.text().strip(),
                            app_version=__version__, **self.meta)

    def _default_name(self, ext):
        p = self.profile()
        stem = p.part_number or "flyback"
        return f"{stem}_{datetime.now():%Y%m%d_%H%M}.{ext}".replace(" ", "_").replace("/", "-")

    def preview(self):
        img = render_image(self.profile(), self.results, self._meta(), dpi=100)
        ReportPreview(img, self.export_pdf, self).exec()

    def export_pdf(self):
        path, _ = QFileDialog.getSaveFileName(self, "Export report", self._default_name("pdf"), "PDF (*.pdf)")
        if path:
            verdict = export_pdf(path, self.profile(), self.results, self._meta())
            self.progress_text.setText(f"Report saved ({verdict}): {path}")

    def export_data(self):
        path, _ = QFileDialog.getSaveFileName(self, "Export raw data", self._default_name("xlsx"),
                                              "Excel workbook (*.xlsx)")
        if not path:
            return
        from openpyxl import Workbook
        from openpyxl.styles import Font
        p = self.profile()
        wb = Workbook()
        ws = wb.active
        ws.title = "Summary"
        meta = self._meta()
        for k in ("date", "station", "meter", "correction", "notes"):
            ws.append([k.title(), meta.get(k, "")])
        ws.append(["Part number", p.part_number])
        ws.append([])
        rows, verdict = evaluate(p, self.results)
        ws.append(["Verdict", verdict])
        ws.append(["Parameter", "Condition", "Measured", "Limit", "Result"])
        for c in ws[ws.max_row]:
            c.font = Font(bold=True)
        for r in rows:
            ws.append([r["param"], r["cond"], r["value"], r["limit"], r["status"]])
        names = {"lp": "Lp (open)", "llk": "Llk (shorted)"}
        for key, res in self.results.items():
            sh = wb.create_sheet(names.get(key, key)[:31])
            equ = res.get("equ")
            sh.append(["Level", "Frequency (Hz)", "L (H)", "L σ (H)", "Q", "D", "Phase (°)",
                       "Rs (Ω)" if equ == "SER" else "Rp (Ω)", "Model"])
            for c in sh[1]:
                c.font = Font(bold=True)
            for r in res["rows"]:
                d, theta, rr = loss_figures(r, equ)
                sh.append([r["level"], r["hz"], r["L"], r["L_sd"], r["Q"], d, theta, rr, equ])
        wb.save(path)
        self.progress_text.setText(f"Data saved: {path}")
